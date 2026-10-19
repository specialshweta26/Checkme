from __future__ import annotations

import io
from functools import lru_cache
from pathlib import Path

import boto3
from botocore.client import Config
from botocore.exceptions import ClientError

from app.config import get_settings


def _disk_root() -> Path:
    root = Path(get_settings().storage_disk_path)
    root.mkdir(parents=True, exist_ok=True)
    return root


def _disk_file(key: str) -> Path:
    dest = (_disk_root() / key).resolve()
    if not str(dest).startswith(str(_disk_root().resolve())):
        raise ValueError("invalid storage key")
    return dest


def guess_media_type(filename: str, content_type: str | None = None) -> str:
    name = (filename or "").lower()
    ctype = (content_type or "").split(";")[0].strip().lower()
    if name.endswith(".pdf") or ctype == "application/pdf":
        return "application/pdf"
    if name.endswith(".png") or ctype == "image/png":
        return "image/png"
    if name.endswith((".jpg", ".jpeg")) or ctype in {"image/jpeg", "image/jpg"}:
        return "image/jpeg"
    if name.endswith(".gif"):
        return "image/gif"
    if name.endswith(".webp"):
        return "image/webp"
    if name.endswith(".txt"):
        return "text/plain"
    if name.endswith(".html"):
        return "text/html"
    return ctype or "application/octet-stream"


def render_pdf_preview(data: bytes, max_pages: int = 2) -> bytes:
    import pypdfium2 as pdfium
    from PIL import Image

    pdf = pdfium.PdfDocument(data)
    images = []
    try:
        for i in range(min(len(pdf), max_pages)):
            page = pdf[i]
            try:
                images.append(page.render(scale=1.7).to_pil().convert("RGB"))
            finally:
                page.close()
    finally:
        pdf.close()
    if not images:
        raise ValueError("empty pdf")
    if len(images) == 1:
        out = images[0]
    else:
        width = max(im.width for im in images)
        gap = 16
        height = sum(im.height for im in images) + gap * (len(images) - 1)
        out = Image.new("RGB", (width, height), (255, 255, 255))
        y = 0
        for im in images:
            out.paste(im, (0, y))
            y += im.height + gap
    buf = io.BytesIO()
    out.save(buf, format="PNG")
    return buf.getvalue()


@lru_cache
def _client():
    s = get_settings()
    endpoint = s.s3_endpoint_url
    if endpoint and "://" not in endpoint:
        endpoint = f"http://{endpoint}"
    return boto3.client(
        "s3",
        endpoint_url=endpoint,
        aws_access_key_id=s.s3_access_key,
        aws_secret_access_key=s.s3_secret_key,
        config=Config(signature_version="s3v4"),
        region_name="us-east-1",
    )


def ensure_bucket() -> None:
    if get_settings().use_disk_storage:
        _disk_root()
        return
    s = get_settings()
    client = _client()
    try:
        client.head_bucket(Bucket=s.s3_bucket)
    except ClientError:
        client.create_bucket(Bucket=s.s3_bucket)


def put_bytes(key: str, data: bytes, content_type: str) -> None:
    if get_settings().use_disk_storage:
        dest = _disk_file(key)
        dest.parent.mkdir(parents=True, exist_ok=True)
        dest.write_bytes(data)
        dest.with_name(dest.name + ".ctype").write_text(content_type or "", encoding="utf-8")
        return
    _client().put_object(Bucket=get_settings().s3_bucket, Key=key, Body=data, ContentType=content_type)


def get_bytes(key: str) -> tuple[bytes, str]:
    if get_settings().use_disk_storage:
        dest = _disk_file(key)
        data = dest.read_bytes()
        meta = dest.with_name(dest.name + ".ctype")
        stored = meta.read_text(encoding="utf-8").strip() if meta.exists() else ""
        return data, stored or guess_media_type(dest.name)
    obj = _client().get_object(Bucket=get_settings().s3_bucket, Key=key)
    data = obj["Body"].read()
    return data, obj.get("ContentType") or "application/octet-stream"


def file_url(key: str, filename: str, content_type: str, inline: bool = False) -> str:
    from urllib.parse import quote

    from app.security import sign_state

    token = sign_state({"key": key, "filename": filename, "content_type": content_type, "inline": inline})
    return f"{get_settings().public_web_url.rstrip('/')}/api/files?t={quote(token, safe='')}"


def signed_url(key: str, filename: str, expires: int = 3600, inline: bool = False, content_type: str = "application/octet-stream") -> str:
    return file_url(key, filename, content_type, inline=inline)


def extract_text(filename: str, content_type: str, data: bytes, limit: int = 4000) -> str:
    name = filename.lower()
    if name.endswith(".txt") or name.endswith(".csv") or name.endswith(".md") or content_type.startswith("text/"):
        return data.decode("utf-8", "replace")[:limit]
    if name.endswith(".html") or content_type == "text/html":
        from connector_core.text import html_to_text

        return html_to_text(data.decode("utf-8", "replace"))[:limit]
    if name.endswith(".pdf") or content_type == "application/pdf":
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(data))
        pages = []
        for page in reader.pages[:8]:
            pages.append(page.extract_text() or "")
        return "\n".join(pages)[:limit]
    if name.endswith(".docx"):
        from docx import Document

        doc = Document(io.BytesIO(data))
        return "\n".join(p.text for p in doc.paragraphs)[:limit]
    return ""
